"""SpreadsheetBench rollout — codegen & ReAct batch execution.

Provides:
  - process_one_codegen(): single/multi-round code generation (no tool-call)
  - run_spreadsheet_batch_codegen(): batch wrapper for codegen
  - process_one(): ReAct agent with tool-call (legacy)
  - run_spreadsheet_batch(): batch wrapper for ReAct (legacy)
  - load_items(): load benchmark .json/.jsonl files
"""
from __future__ import annotations

import glob as _glob
import json
import os
import shutil
import tempfile
import time
import traceback
from concurrent.futures import (
    FIRST_COMPLETED,
    ThreadPoolExecutor,
    wait,
)

import openpyxl

from skillopt.envs.spreadsheetbench.react_agent import run_react
from skillopt.envs.spreadsheetbench.evaluator import (
    evaluate, _generate_cell_names, _compare_cell_value,
)
from skillopt.envs.spreadsheetbench.executor import run_generated_code
from skillopt.model.infra_errors import InfraError, classify_infra_error
from skillopt.engine.run_artifacts import assert_valid_results, write_json, write_stage_stats


def _failure_type(result: dict) -> str:
    if result.get("failure_type"):
        return result["failure_type"]
    if result.get("ok"):
        return "none"
    reason = result.get("fail_reason", "")
    if "timeout" in reason.lower():
        return "task_timeout"
    return {"setup": "dataset_error", "llm": "llm_output_error",
            "agent": "agent_error", "extract": "code_missing",
            "exec": "execution_error" if not result.get("exec_ok") else "score_mismatch",
            "error": "unexpected_error"}.get(result.get("phase"), "score_mismatch")


def _persist_task_result(out_root: str, result: dict, started_at: float) -> None:
    """Persist even an attempt that obtained no assistant response."""
    task_out = os.path.join(out_root, "predictions", str(result["id"]))
    os.makedirs(task_out, exist_ok=True)
    result.setdefault("status", "completed")
    result["failure_type"] = _failure_type(result)
    result["wall_time_s"] = round(time.time() - started_at, 3)
    conversation_path = os.path.join(task_out, "conversation.json")
    if not os.path.exists(conversation_path):
        # An empty observed transcript is truthful. Do not invent an assistant
        # answer or make the infrastructure error look like a model message.
        write_json(conversation_path, [])
    with open(conversation_path, encoding="utf-8") as handle:
        conversation = json.load(handle)
    write_json(os.path.join(task_out, "conversation_meta.json"), {
        "status": result["status"], "partial": result["status"] == "infra_error",
        "assistant_response_observed": any(
            m.get("role") == "assistant" and bool(m.get("content")) for m in conversation
        ),
        "failure_type": result["failure_type"],
        "note": "Only observed messages are stored; an empty list means no response was captured.",
    })
    raw_path = os.path.join(task_out, "raw.txt")
    if not os.path.exists(raw_path):
        with open(raw_path, "w", encoding="utf-8") as handle:
            handle.write("")
    trace_path = os.path.join(task_out, "raw_trace.txt")
    if not os.path.exists(trace_path):
        evidence_path = os.path.join(task_out, "infra_raw_trace.txt")
        if not os.path.exists(evidence_path):
            evidence_path = raw_path
        with open(evidence_path, encoding="utf-8") as handle:
            raw = handle.read()
        with open(trace_path, "w", encoding="utf-8") as handle:
            handle.write(raw or result.get("error", ""))
    result["conversation_path"] = conversation_path
    result["raw_trace_path"] = trace_path
    write_json(os.path.join(task_out, "result.json"), result)
    write_stage_stats(task_out, [result], expected=1,
                      status=result["status"], started_at=started_at)


def _raise_task_infra(result: dict, exc: Exception, *, stage: str,
                      model_call: bool = False) -> None:
    # Workbook/cell errors can contain numbers such as 401 and local execution
    # can time out. Only actual model-call boundaries interpret legacy text.
    infra = exc if isinstance(exc, InfraError) else (
        classify_infra_error(exc, stage=stage, role="target") if model_call else None
    )
    if infra is None:
        return
    result.update(status="infra_error", failure_type=infra.failure_type,
                  infra_error=infra.to_dict(), hard=None, soft=None, ok=False,
                  fail_reason=str(infra), error=json.dumps(infra.to_dict(), ensure_ascii=False))
    infra.details["task_id"] = result["id"]
    infra.task_result = result
    raise infra from exc


def _append_result(results_path: str, result: dict) -> None:
    with open(results_path, "a", encoding="utf-8") as handle:
        handle.write(json.dumps(result, ensure_ascii=False) + "\n")
        handle.flush()


def _batch_infra_result(item: dict, error: InfraError) -> dict:
    return getattr(error, "task_result", None) or {
        "id": str(item["id"]), "status": "infra_error", "ok": False,
        "hard": None, "soft": None, "phase": error.stage,
        "failure_type": error.failure_type, "infra_error": error.to_dict(),
        "fail_reason": str(error), "llm_ok": False, "code_ok": False,
        "exec_ok": False, "n_turns": 0,
    }


def _run_task_batch(items, out_root, run_one, *, workers, task_timeout) -> list[dict]:
    """Bound submission and persist each outcome before dispatching another."""
    if workers < 1:
        raise ValueError("workers must be >= 1")
    os.makedirs(out_root, exist_ok=True)
    started = time.time()
    results_path = os.path.join(out_root, "results.jsonl")
    requested_ids = {str(item["id"]) for item in items}
    results = []
    if os.path.exists(results_path):
        with open(results_path, encoding="utf-8") as handle:
            for line in handle:
                if line.strip():
                    row = json.loads(line)
                    if str(row["id"]) in requested_ids:
                        results.append(row)
    assert_valid_results(results, stage=out_root)
    done_ids = {str(row["id"]) for row in results}
    pending = [item for item in items if str(item["id"]) not in done_ids]
    state = "running"
    write_stage_stats(out_root, results, expected=len(items), status=state, started_at=started)

    def record(row):
        row.setdefault("status", "completed")
        row["failure_type"] = _failure_type(row)
        results.append(row)
        _append_result(results_path, row)
        write_stage_stats(out_root, results, expected=len(items), status=state, started_at=started)
        print(f"    {len(results)}/{len(items)} id={row['id']} "
              f"status={row['status']} failure={row['failure_type']} "
              f"hard={row.get('hard')} elapsed={time.time() - started:.0f}s", flush=True)

    def fail(item, error):
        nonlocal state
        state = "infra_error"
        row = _batch_infra_result(item, error)
        _persist_task_result(out_root, row, started)
        record(row)
        write_json(os.path.join(out_root, "infra_error.json"), error.to_dict())
        raise error

    # Serial dispatch guarantees a failing first model request never launches
    # task two. The model backend owns the hard subprocess/request deadline.
    if workers == 1:
        for item in pending:
            try:
                row = run_one(item)
                assert_valid_results([row], stage=out_root)
            except Exception as exc:
                infra = exc if isinstance(exc, InfraError) else None
                if infra is not None:
                    fail(item, infra)
                raise
            record(row)
    else:
        iterator = iter(pending)
        executor = ThreadPoolExecutor(max_workers=workers)
        active = {}
        def submit_one():
            item = next(iterator, None)
            if item is not None:
                active[executor.submit(run_one, item)] = (item, time.monotonic())
        try:
            for _ in range(min(workers, len(pending))):
                submit_one()
            while active:
                done, _ = wait(active, timeout=1, return_when=FIRST_COMPLETED)
                for future in done:
                    item, _ = active.pop(future)
                    try:
                        row = future.result()
                        assert_valid_results([row], stage=out_root)
                    except Exception as exc:
                        infra = exc if isinstance(exc, InfraError) else None
                        if infra is not None:
                            fail(item, infra)
                        raise
                    record(row)
                for future, (item, began) in active.items():
                    if task_timeout > 0 and time.monotonic() - began >= task_timeout:
                        fail(item, InfraError("worker_timeout", "Worker exceeded task deadline; batch aborted",
                                              stage=out_root, details={"timeout_s": task_timeout}))
                for _ in range(workers - len(active)):
                    submit_one()
        finally:
            for future in active:
                future.cancel()
            executor.shutdown(wait=False, cancel_futures=True)
    write_stage_stats(out_root, results, expected=len(items), status="completed", started_at=started)
    return results


# ── Data loading ─────────────────────────────────────────────────────────────


def load_items(path: str) -> list[dict]:
    """Load a benchmark file. Supports both .jsonl and .json (list of dicts)."""
    if path.endswith(".json"):
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            data = data.get("data") or list(data.values())
        return list(data)
    items = []
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                items.append(json.loads(line))
    return items


# ── Test case discovery ──────────────────────────────────────────────────────


def _find_test_cases(task_dir: str) -> list[tuple[str, str, str]]:
    """Return [(case_no, input_path, answer_path), ...] sorted by case_no.

    Supports naming conventions used by SpreadsheetBench releases:
      * ``{no}_{id}_input.xlsx``  + ``{no}_{id}_answer.xlsx``  (original)
      * ``{no}_{id}_init.xlsx``   + ``{no}_{id}_golden.xlsx``  (verified_400)
      * ``initial.xlsx``          + ``golden.xlsx``             (verified_400, no prefix)
    """
    cases: list[tuple[str, str, str]] = []
    inputs = sorted(_glob.glob(os.path.join(task_dir, "*_input.xlsx")))
    for ip in inputs:
        no = os.path.basename(ip).split("_", 1)[0]
        ap = ip.replace("_input.xlsx", "_answer.xlsx")
        if os.path.exists(ap):
            cases.append((no, ip, ap))
    inits = sorted(_glob.glob(os.path.join(task_dir, "*_init.xlsx")))
    for ip in inits:
        no = os.path.basename(ip).split("_", 1)[0]
        ap = ip.replace("_init.xlsx", "_golden.xlsx")
        if os.path.exists(ap):
            cases.append((no, ip, ap))

    if not cases:
        golden_files = sorted(_glob.glob(os.path.join(task_dir, "*_golden.xlsx")))
        if len(inits) == 1 and len(golden_files) == 1:
            no = os.path.basename(inits[0]).split("_", 1)[0]
            cases.append((no, inits[0], golden_files[0]))

    # Fallback: bare initial.xlsx + golden.xlsx (no numbered prefix)
    if not cases:
        bare_init = os.path.join(task_dir, "initial.xlsx")
        bare_gold = os.path.join(task_dir, "golden.xlsx")
        if os.path.exists(bare_init) and os.path.exists(bare_gold):
            cases.append(("1", bare_init, bare_gold))

    return cases


# ── Auto-verify helper ──────────────────────────────────────────────────────

# The official SpreadsheetBench evaluator never serialises cells to text — it
# compares in memory and returns only a pass/fail bool. The per-cell report
# below is a repo-local training aid (fed back to the model on retry and saved
# into the trajectory for reflection). On most tasks the answer range is a
# handful of cells, so the full report is tiny. But a few tasks have answer
# ranges spanning tens of thousands of cells (e.g. 80-42 =
# 'Consolidate_ALL'!A2:L8000 ≈ 96k cells); dumping every cell explodes the
# report to several MB, floods the model's context and bloats conversation
# files. We therefore apply the same head+tail character truncation the rest of
# the codebase uses for oversized trajectory text (cf. reflect.py / slow_update.py
# `text[:half] + "...[truncated]...\n" + text[-half:]`): keep the first and last
# `_MAX_REPORT_CHARS // 2` chars so both the leading and trailing wrong cells
# stay visible. Small reports are unchanged.
_MAX_REPORT_CHARS = 12000      # head+tail char budget (~6000 head + 6000 tail)


def _auto_verify_output(
    pred_path: str,
    gold_path: str,
    answer_position: str,
) -> str:
    """Reopen the predicted xlsx and compare cells at answer_position with gold.

    Returns a human-readable verification report that can be appended to the
    trajectory so the error analyst can see exactly what went wrong (e.g.
    ``cell A1: got=None, expected=420``). Oversized reports are head+tail
    truncated to `_MAX_REPORT_CHARS` chars, matching the rest of the codebase.
    """
    if not os.path.exists(pred_path):
        return "Verification: output file does not exist."
    try:
        wb_pred = openpyxl.load_workbook(pred_path, data_only=True)
        wb_gold = openpyxl.load_workbook(gold_path, data_only=True)
    except Exception as e:
        return f"Verification: could not open workbooks: {e}"

    lines = ["## Output Verification"]
    try:
        for scr in (answer_position or "").split(","):
            scr = scr.strip()
            if not scr:
                continue
            if "!" in scr:
                sheet_name, cell_range = scr.split("!", 1)
                sheet_name = sheet_name.strip().strip("'\"")
            else:
                sheet_name = wb_gold.sheetnames[0]
                cell_range = scr
            cell_range = cell_range.strip().strip("'\"")

            cell_names = _generate_cell_names(cell_range)
            ws_pred = wb_pred[sheet_name] if sheet_name in wb_pred.sheetnames else None
            ws_gold = wb_gold[sheet_name] if sheet_name in wb_gold.sheetnames else None

            if ws_pred is None:
                lines.append(f"  Sheet '{sheet_name}' NOT FOUND in output.")
                continue

            n_empty_correct = 0   # empty-on-both correct cells collapsed to a count
            for cn in cell_names:
                gv = ws_gold[cn].value if ws_gold else "N/A"
                pv = ws_pred[cn].value
                # Use the official cell comparator so this report's ✓/✗ agrees
                # with the real scorer (evaluate). repr() equality would wrongly
                # flag e.g. 5 vs 5.0 or None vs "" as mismatches and mislead the
                # model into "fixing" cells that already pass scoring.
                ok_cell = ws_gold is not None and _compare_cell_value(gv, pv)
                # Collapse only cells that are correct AND empty on both sides
                # (got=None, expected=None ✓): pure noise. Every other cell —
                # including non-empty correct cells — is listed in full; the
                # final head+tail char cap keeps the report bounded.
                if ok_cell and gv in (None, "") and pv in (None, ""):
                    n_empty_correct += 1
                    continue
                match = "✓" if ok_cell else "✗"
                lines.append(f"  {sheet_name}!{cn}: got={pv!r}, expected={gv!r} {match}")
            if n_empty_correct:
                lines.append(
                    f"  (+{n_empty_correct} empty cells correct, omitted)"
                )

        # Also check if any cells in the output contain formula strings
        formula_cells = []
        for sn in wb_pred.sheetnames:
            ws = wb_pred[sn]
            for row in ws.iter_rows(max_row=min(ws.max_row, 200), values_only=False):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formula_cells.append(f"{sn}!{cell.coordinate}={cell.value}")
                        if len(formula_cells) >= 10:
                            break
                if len(formula_cells) >= 10:
                    break
            if len(formula_cells) >= 10:
                break
        if formula_cells:
            lines.append(f"\n  WARNING: {len(formula_cells)} cells contain Excel formulas (openpyxl cannot evaluate them):")
            for fc in formula_cells[:5]:
                lines.append(f"    {fc}")
            if len(formula_cells) > 5:
                lines.append(f"    ... and {len(formula_cells) - 5} more")
    finally:
        wb_pred.close()
        wb_gold.close()

    report = "\n".join(lines)
    # Head+tail truncation, matching reflect.py / slow_update.py: keep the first
    # and last half so both leading and trailing wrong cells remain visible.
    if len(report) > _MAX_REPORT_CHARS:
        half = _MAX_REPORT_CHARS // 2
        report = (
            report[:half]
            + f"\n  ...[verification report truncated, {len(report)} chars total]...\n"
            + report[-half:]
        )
    return report


# ── Per-task worker ──────────────────────────────────────────────────────────


def process_one(
    item: dict,
    data_root: str,
    out_root: str,
    skill_content: str,
    max_turns: int,
    diagnostic_mode: bool = False,
    diagnostic_instruction: str = "",
    diagnostic_trace_context: str = "",
    max_completion_tokens: int = 16384,
) -> dict:
    """Run the ReAct agent on a single SpreadsheetBench task.

    Returns a result dict compatible with ``compute_score()``.
    """
    task_id = str(item["id"])
    instruction = item["instruction"]
    instruction_type = item.get("instruction_type", "")
    answer_position = item.get("answer_position", "")
    answer_sheet = item.get("answer_sheet", "")
    if answer_position and answer_sheet and "!" not in answer_position:
        answer_position_eval = f"{answer_sheet}!{answer_position}"
    else:
        answer_position_eval = answer_position

    # Determine task_type from instruction_type
    itype_lower = (instruction_type or "").lower()
    if "cell" in itype_lower:
        task_type = "cell_level"
    elif "sheet" in itype_lower:
        task_type = "sheet_level"
    else:
        task_type = "other"

    sp = item.get("spreadsheet_path", f"spreadsheet/{task_id}")
    task_dir = sp if os.path.isabs(sp) else os.path.join(data_root, sp)
    if not os.path.exists(task_dir):
        abs_task_dir = os.path.abspath(task_dir)
        if os.path.exists(abs_task_dir):
            task_dir = abs_task_dir

    result = {
        "id": task_id,
        "ok": False,
        "instruction_type": instruction_type,
        "task_type": task_type,
        "task_description": instruction,
        "phase": "setup",
        "fail_reason": "",
        "agent_ok": False,
        "exec_ok": False,
        "n_cases": 0,
        "n_exec_pass": 0,
        "n_pass": 0,
        "soft": 0.0,
        "hard": 0,
        "n_turns": 0,
        "cases": [],
        "error": "",
    }

    task_started_at = time.time()
    try:
        cases = _find_test_cases(task_dir)
        result["n_cases"] = len(cases)
        if not cases:
            result["fail_reason"] = "no-test-cases"
            result["error"] = (
                f"task_dir={os.path.abspath(task_dir)} "
                f"exists={os.path.exists(task_dir)} cwd={os.getcwd()} "
                f"data_root_repr={data_root!r} spreadsheet_path_repr={sp!r} task_dir_repr={task_dir!r}"
            )
            return result

        task_out_dir = os.path.join(out_root, "predictions", task_id)
        os.makedirs(task_out_dir, exist_ok=True)

        no1, ip1, _ = cases[0]
        pred_path_1 = os.path.join(task_out_dir, f"{no1}_pred.xlsx")
        target_prompt_parts = [
            f"# Instruction\n{instruction}",
            f"# Input file\n{ip1}",
            f"# Output file\n{pred_path_1}",
        ]
        if instruction_type:
            target_prompt_parts.append(f"# Instruction type\n{instruction_type}")
        if answer_position_eval:
            target_prompt_parts.append(f"# Answer position\n{answer_position_eval}")
        if diagnostic_trace_context.strip():
            target_prompt_parts.insert(
                0,
                "# Previous Codex Trace Snapshot\n"
                "This is a partial transcript from an earlier attempt. Use it as your current reasoning context.\n\n"
                f"{diagnostic_trace_context.strip()}",
            )
        if diagnostic_mode and diagnostic_instruction.strip():
            target_prompt_parts.append(f"# Training readout\n{diagnostic_instruction.strip()}")
        target_user_prompt = "\n\n".join(target_prompt_parts)
        try:
            from skillopt.envs.spreadsheetbench.react_agent import _build_system
            target_system_prompt = _build_system(skill_content)
        except Exception:
            target_system_prompt = ""
        if target_system_prompt:
            with open(os.path.join(task_out_dir, "target_system_prompt.txt"), "w", encoding="utf-8") as f:
                f.write(target_system_prompt)
            result["target_system_prompt"] = target_system_prompt
        with open(os.path.join(task_out_dir, "target_user_prompt.txt"), "w", encoding="utf-8") as f:
            f.write(target_user_prompt)
        result["target_user_prompt"] = target_user_prompt

        # ── Stage 1: run ReAct agent on test case 1 ─────────────────────
        result["phase"] = "agent"

        work_dir = tempfile.mkdtemp(prefix=f"react_{task_id}_")
        try:
            # Copy input so agent works in an isolated directory
            work_input = os.path.join(work_dir, os.path.basename(ip1))
            shutil.copy2(ip1, work_input)

            agent_result = run_react(
                instruction=instruction,
                input_path=work_input,
                output_path=pred_path_1,
                work_dir=work_dir,
                instruction_type=instruction_type,
                answer_position=answer_position_eval,
                skill_content=skill_content,
                max_turns=max_turns,
                max_output_tokens=max_completion_tokens,
                diagnostic_mode=diagnostic_mode,
                diagnostic_instruction=diagnostic_instruction,
                diagnostic_trace_context=diagnostic_trace_context,
            )
            result["n_turns"] = agent_result.get("n_turns", 0)
            if agent_result.get("target_system_prompt"):
                with open(os.path.join(task_out_dir, "target_system_prompt.txt"), "w", encoding="utf-8") as f:
                    f.write(agent_result["target_system_prompt"])
                result["target_system_prompt"] = agent_result["target_system_prompt"]
            if agent_result.get("target_user_prompt"):
                with open(os.path.join(task_out_dir, "target_user_prompt.txt"), "w", encoding="utf-8") as f:
                    f.write(agent_result["target_user_prompt"])
                result["target_user_prompt"] = agent_result["target_user_prompt"]

            # Save conversation log
            with open(os.path.join(task_out_dir, "conversation.json"), "w", encoding="utf-8") as f:
                json.dump(
                    agent_result.get("conversation", []),
                    f, ensure_ascii=False, indent=2,
                )

            # Copy solution.py if the agent wrote one
            solution_src = os.path.join(work_dir, "solution.py")
            solution_dst = os.path.join(task_out_dir, "solution.py")
            if os.path.exists(solution_src):
                shutil.copy2(solution_src, solution_dst)

        except Exception as e:
            _raise_task_infra(result, e, stage="target_react", model_call=True)
            result["fail_reason"] = f"agent-error: {type(e).__name__}: {e}"
            result["error"] = traceback.format_exc()
            return result
        finally:
            shutil.rmtree(work_dir, ignore_errors=True)

        result["agent_ok"] = True

        # ── Stage 2: evaluate all test cases ─────────────────────────────
        result["phase"] = "eval"
        solution_path = os.path.join(task_out_dir, "solution.py")
        all_exec = True

        for i, (no, ip, ap) in enumerate(cases):
            pred_path = os.path.join(task_out_dir, f"{no}_pred.xlsx")

            if i > 0:
                # Re-apply solution.py to subsequent test cases
                if not os.path.exists(solution_path):
                    all_exec = False
                    result["cases"].append(
                        {"no": no, "stage": "exec", "ok": False, "error": "no-solution-py"}
                    )
                    if not result["fail_reason"]:
                        result["fail_reason"] = "no-solution-py-for-other-cases"
                    continue

                with open(solution_path, encoding="utf-8") as f:
                    code = f.read()

                # Prepend new INPUT_PATH / OUTPUT_PATH
                preamble = (
                    f"INPUT_PATH  = {ip!r}\n"
                    f"OUTPUT_PATH = {pred_path!r}\n"
                )
                full_code = preamble + code

                ok_exec, err = run_generated_code(full_code, ip, pred_path)
                if not ok_exec:
                    all_exec = False
                    result["cases"].append(
                        {"no": no, "stage": "exec", "ok": False, "error": err[:500]}
                    )
                    if not result["fail_reason"]:
                        tail = err.strip().splitlines()[-1][:200] if err.strip() else "unknown"
                        result["fail_reason"] = f"exec-error: {tail}"
                    continue

            # ── Evaluate ─────────────────────────────────────────────────
            if not os.path.exists(pred_path):
                all_exec = False
                result["cases"].append(
                    {"no": no, "stage": "exec", "ok": False, "error": "output-not-found"}
                )
                if not result["fail_reason"]:
                    result["fail_reason"] = "output-not-found"
                continue

            result["n_exec_pass"] += 1
            try:
                ev = evaluate(pred_path, ap, instruction_type, answer_position_eval)
            except Exception as e:  # noqa: BLE001
                ev = {"ok": False, "reason": f"eval-exception: {type(e).__name__}: {e}"}

            if ev["ok"]:
                result["n_pass"] += 1
            else:
                if not result["fail_reason"]:
                    result["fail_reason"] = f"eval-mismatch: {ev['reason'][:200]}"
            result["cases"].append(
                {"no": no, "stage": "eval", "ok": ev["ok"], "reason": ev.get("reason", "")}
            )

        result["exec_ok"] = all_exec
        n_cases = result["n_cases"]
        n_pass = result["n_pass"]
        result["soft"] = (n_pass / n_cases) if n_cases else 0.0
        result["hard"] = 1 if (n_cases > 0 and n_pass == n_cases) else 0
        result["ok"] = bool(result["hard"])
        if result["ok"]:
            result["fail_reason"] = ""
        return result

    except Exception as e:  # noqa: BLE001
        _raise_task_infra(result, e, stage=result["phase"])
        result["fail_reason"] = f"unexpected: {type(e).__name__}: {e}"
        result["error"] = traceback.format_exc()
        return result
    finally:
        _persist_task_result(out_root, result, task_started_at)


# ── Batch runner ─────────────────────────────────────────────────────────────


def run_spreadsheet_batch(
    items: list[dict],
    data_root: str,
    out_root: str,
    skill_content: str,
    max_turns: int = 30,
    max_completion_tokens: int = 16384,
    max_api_workers: int = 64,
    task_timeout: int = 600,
    diagnostic_mode: bool = False,
    diagnostic_instruction: str = "",
    diagnostic_trace_context_by_id: dict[str, str] | None = None,
) -> list[dict]:
    """Run legacy ReAct tasks with persistent evidence and fail-fast infra errors."""
    def run_one(item):
        return process_one(
            item, data_root, out_root, skill_content, max_turns,
            diagnostic_mode, diagnostic_instruction,
            (diagnostic_trace_context_by_id or {}).get(str(item["id"]), ""),
            max_completion_tokens,
        )
    return _run_task_batch(items, out_root, run_one, workers=max_api_workers,
                           task_timeout=task_timeout)


# ── Codegen per-task worker (no tool-call) ──────────────────────────────────


def process_one_codegen(
    item: dict,
    data_root: str,
    out_root: str,
    skill_content: str,
    mode: str = "single",
    max_turns: int = 5,
    max_completion_tokens: int = 16384,
    task_timeout: int = 600,
    use_eval_feedback: bool = False,
    diagnostic_mode: bool = False,
    diagnostic_instruction: str = "",
    diagnostic_trace_context: str = "",
) -> dict:
    """Run codegen agent (single or multi-round) on one SpreadsheetBench task.

    This matches the official evaluation setting: LLM generates a Python code
    block, no function-calling / tool-use.
    """
    from skillopt.envs.spreadsheetbench.codegen_agent import run_single, run_multi

    task_id = str(item["id"])
    instruction = item["instruction"]
    instruction_type = item.get("instruction_type", "")
    answer_position = item.get("answer_position", "")
    answer_sheet = item.get("answer_sheet", "")
    if answer_position and answer_sheet and "!" not in answer_position:
        answer_position_eval = f"{answer_sheet}!{answer_position}"
    else:
        answer_position_eval = answer_position

    itype_lower = (instruction_type or "").lower()
    if "cell" in itype_lower:
        task_type = "cell_level"
    elif "sheet" in itype_lower:
        task_type = "sheet_level"
    else:
        task_type = "other"

    sp = item.get("spreadsheet_path", f"spreadsheet/{task_id}")
    task_dir = sp if os.path.isabs(sp) else os.path.join(data_root, sp)
    if not os.path.exists(task_dir):
        abs_task_dir = os.path.abspath(task_dir)
        if os.path.exists(abs_task_dir):
            task_dir = abs_task_dir

    result = {
        "id": task_id,
        "ok": False,
        "instruction_type": instruction_type,
        "task_type": task_type,
        "task_description": instruction,
        "phase": "setup",
        "fail_reason": "",
        "llm_ok": False,
        "code_ok": False,
        "exec_ok": False,
        "n_cases": 0,
        "n_exec_pass": 0,
        "n_pass": 0,
        "soft": 0.0,
        "hard": 0,
        "n_turns": 0,
        "cases": [],
        "error": "",
    }

    task_started_at = time.time()
    try:
        cases = _find_test_cases(task_dir)
        result["n_cases"] = len(cases)
        if not cases:
            result["fail_reason"] = "no-test-cases"
            result["error"] = (
                f"task_dir={os.path.abspath(task_dir)} "
                f"exists={os.path.exists(task_dir)} cwd={os.getcwd()} "
                f"data_root_repr={data_root!r} spreadsheet_path_repr={sp!r} task_dir_repr={task_dir!r}"
            )
            return result

        task_out_dir = os.path.join(out_root, "predictions", task_id)
        os.makedirs(task_out_dir, exist_ok=True)

        # ── Save context for Optimizer (Reflect stage) ──────────────────
        from skillopt.envs.spreadsheetbench.codegen_agent import (
            _preview_workbook, _build_system, _build_user,
        )
        first_input_for_preview = cases[0][1]
        try:
            preview_text = _preview_workbook(first_input_for_preview)
        except Exception:
            preview_text = "(preview failed)"
        target_system = _build_system(skill_content)
        target_user = _build_user(
            instruction,
            first_input_for_preview,
            instruction_type,
            answer_position_eval,
            diagnostic_mode=diagnostic_mode,
            diagnostic_instruction=diagnostic_instruction,
            diagnostic_trace_context=diagnostic_trace_context,
        )

        with open(os.path.join(task_out_dir, "spreadsheet_preview.txt"), "w", encoding="utf-8") as f:
            f.write(preview_text)
        with open(os.path.join(task_out_dir, "target_system_prompt.txt"), "w", encoding="utf-8") as f:
            f.write(target_system)
        with open(os.path.join(task_out_dir, "target_user_prompt.txt"), "w", encoding="utf-8") as f:
            f.write(target_user)

        result["spreadsheet_preview"] = preview_text
        result["target_system_prompt"] = target_system
        result["target_user_prompt"] = target_user

        # ── LLM phase ──────────────────────────────────────────────────
        result["phase"] = "llm"
        first_input = cases[0][1]
        first_gold = cases[0][2]
        first_pred = os.path.join(task_out_dir, f"{cases[0][0]}_pred.xlsx")

        try:
            if mode == "multi":
                agent_result = run_multi(
                    instruction=instruction,
                    input_xlsx=first_input,
                    output_path=first_pred,
                    instruction_type=instruction_type,
                    answer_position=answer_position_eval,
                    skill_content=skill_content,
                    max_turns=max_turns,
                    max_output_tokens=max_completion_tokens,
                    task_timeout=task_timeout,
                    gold_path=first_gold if use_eval_feedback else "",
                    diagnostic_mode=diagnostic_mode,
                    diagnostic_instruction=diagnostic_instruction,
                    diagnostic_trace_context=diagnostic_trace_context,
                )
            else:
                agent_result = run_single(
                    instruction=instruction,
                    input_xlsx=first_input,
                    output_path=first_pred,
                    instruction_type=instruction_type,
                    answer_position=answer_position_eval,
                    skill_content=skill_content,
                    max_output_tokens=max_completion_tokens,
                    task_timeout=task_timeout,
                    diagnostic_mode=diagnostic_mode,
                    diagnostic_instruction=diagnostic_instruction,
                    diagnostic_trace_context=diagnostic_trace_context,
                )
        except Exception as e:  # noqa: BLE001
            _raise_task_infra(result, e, stage="target_codegen", model_call=True)
            result["fail_reason"] = f"llm-call-failed: {type(e).__name__}: {e}"
            result["error"] = traceback.format_exc()
            return result

        result["llm_ok"] = True
        result["n_turns"] = agent_result.get("n_turns", 1)
        code = agent_result.get("code", "")
        raw = agent_result.get("raw", "")

        # Save artifacts
        with open(os.path.join(task_out_dir, "code.py"), "w", encoding="utf-8") as f:
            f.write(code)
        with open(os.path.join(task_out_dir, "raw.txt"), "w", encoding="utf-8") as f:
            f.write(raw)
        with open(os.path.join(task_out_dir, "conversation.json"), "w", encoding="utf-8") as f:
            json.dump(agent_result.get("conversation", []), f, ensure_ascii=False, indent=2)

        if not code.strip():
            result["phase"] = "extract"
            result["fail_reason"] = "empty-code-block"
            return result
        try:
            compile(code, "generated_solution.py", "exec")
        except SyntaxError as error:
            result.update(phase="extract", failure_type="code_syntax_error", fail_reason=str(error))
            return result
        result["code_ok"] = True

        # ── Exec + eval per test case ──────────────────────────────────
        result["phase"] = "exec"
        all_exec = True
        # Collect enrichment info for the conversation/trajectory
        enrichment_parts: list[str] = []

        for no, ip, ap in cases:
            pred_path = os.path.join(task_out_dir, f"{no}_pred.xlsx")

            # For multi mode, the first case may already be produced
            if not os.path.exists(pred_path):
                ok_exec, err = run_generated_code(code, ip, pred_path)
                if not ok_exec:
                    all_exec = False
                    result["cases"].append(
                        {"no": no, "stage": "exec", "ok": False, "error": err[:500]}
                    )
                    if not result["fail_reason"]:
                        tail = err.strip().splitlines()[-1][:200] if err.strip() else "unknown"
                        result["fail_reason"] = f"exec-error: {tail}"
                    enrichment_parts.append(
                        f"## Execution (case {no})\nERROR: {err[:500]}"
                    )
                    continue

            if not os.path.exists(pred_path):
                all_exec = False
                result["cases"].append(
                    {"no": no, "stage": "exec", "ok": False, "error": "output-not-found"}
                )
                if not result["fail_reason"]:
                    result["fail_reason"] = "output-not-found"
                continue

            result["n_exec_pass"] += 1
            try:
                ev = evaluate(pred_path, ap, instruction_type, answer_position_eval)
            except Exception as e:  # noqa: BLE001
                ev = {"ok": False, "reason": f"eval-exception: {type(e).__name__}: {e}"}

            if ev["ok"]:
                result["n_pass"] += 1
            else:
                if not result["fail_reason"]:
                    result["fail_reason"] = f"eval-mismatch: {ev['reason'][:200]}"
            result["cases"].append(
                {"no": no, "stage": "eval", "ok": ev["ok"], "reason": ev.get("reason", "")}
            )

            # Auto-verify: reopen output and compare cells at answer_position
            if answer_position_eval:
                verify_report = _auto_verify_output(pred_path, ap, answer_position_eval)
                enrichment_parts.append(
                    f"## Eval Result (case {no}): {'PASS' if ev['ok'] else 'FAIL'}\n"
                    f"{ev.get('reason', '')}\n\n{verify_report}"
                )

        result["exec_ok"] = all_exec

        # ── Enrich conversation with eval details ──────────────────────
        if enrichment_parts:
            enrichment_msg = "\n\n---\n\n".join(enrichment_parts)
            conversation = agent_result.get("conversation", [])
            conversation.append({
                "role": "system",
                "content": f"[POST-EXECUTION VERIFICATION]\n\n{enrichment_msg}",
            })
            # Re-save the enriched conversation
            with open(os.path.join(task_out_dir, "conversation.json"), "w", encoding="utf-8") as f:
                json.dump(conversation, f, ensure_ascii=False, indent=2)
        n_cases = result["n_cases"]
        n_pass = result["n_pass"]
        result["soft"] = (n_pass / n_cases) if n_cases else 0.0
        result["hard"] = 1 if (n_cases > 0 and n_pass == n_cases) else 0
        result["ok"] = bool(result["hard"])
        if result["ok"]:
            result["fail_reason"] = ""
        return result

    except Exception as e:  # noqa: BLE001
        _raise_task_infra(result, e, stage=result["phase"])
        result["fail_reason"] = f"unexpected: {type(e).__name__}: {e}"
        result["error"] = traceback.format_exc()
        return result
    finally:
        _persist_task_result(out_root, result, task_started_at)


# ── Codegen batch runner ────────────────────────────────────────────────────


def run_spreadsheet_batch_codegen(
    items: list[dict],
    data_root: str,
    out_root: str,
    skill_content: str,
    mode: str = "single",
    max_turns: int = 5,
    max_completion_tokens: int = 16384,
    max_api_workers: int = 32,
    task_timeout: int = 0,
    use_eval_feedback: bool = False,
    diagnostic_mode: bool = False,
    diagnostic_instruction: str = "",
    diagnostic_trace_context_by_id: dict[str, str] | None = None,
) -> list[dict]:
    """Run code-generation tasks, aborting the stage on infrastructure errors."""
    def run_one(item):
        return process_one_codegen(
            item=item, data_root=data_root, out_root=out_root,
            skill_content=skill_content, mode=mode, max_turns=max_turns,
            max_completion_tokens=max_completion_tokens, task_timeout=task_timeout,
            use_eval_feedback=use_eval_feedback, diagnostic_mode=diagnostic_mode,
            diagnostic_instruction=diagnostic_instruction,
            diagnostic_trace_context=(diagnostic_trace_context_by_id or {}).get(str(item["id"]), ""),
        )
    return _run_task_batch(items, out_root, run_one, workers=max_api_workers,
                           task_timeout=task_timeout)
